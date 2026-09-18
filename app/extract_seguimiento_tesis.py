"""Bronze -> Silver

Extrae la información oficial de seguimiento de trabajos de grado de los
estudiantes de Investigación (595) y Profundización (695) a partir de las bases
de datos institucionales en Bronze:
  1. Data/Bronze/MCIC - Base de datos INVESTIGACION.xlsx
  2. Data/Bronze/MCIC - Base de datos Profundizacion.xlsx
  3. Data/Bronze/MCIC - Base de datos V2.xlsx

Extrae los campos:
  - cod_proyecto: '595' (Investigación) o '695' (Profundización)
  - etapa_actual: Sustentado (Aprobado/Meritorio/Laureado), Jurados solicitados
    (pendiente sustentación), o Anteproyecto radicado (pendiente jurados).
  - titulo_proyecto
  - director
  - codirector
  - grupo_investigacion
  - numero_acta_sustentacion
  - nota_sustentacion
  - caracter_sustentacion
  - jurados_sustentacion
  - fuente_categorias

Política de privacidad (no negociable): el código y el nombre del estudiante
se usan SOLO en memoria para asociar, normalizar y desduplicar registros.
Nunca se escriben en Data/Silver/seguimiento_tesis.json para garantizar
la protección de datos personales.
"""
from __future__ import annotations

import csv
import datetime
import json
import re
import sys
import unicodedata
import warnings
from pathlib import Path
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SILVER_DIR = ROOT / "Data/Silver"
BRONZE_DIR = ROOT / "Data/Bronze"
ESTADOS_DIR = BRONZE_DIR / "Estados"

ARCHIVO_INV = BRONZE_DIR / "MCIC - Base de datos INVESTIGACION.xlsx"
ARCHIVO_PROF = BRONZE_DIR / "MCIC - Base de datos Profundizacion.xlsx"
ARCHIVO_V2 = BRONZE_DIR / "MCIC - Base de datos V2.xlsx"

COL_COD_ESTUDIANTE = 3
PROYECTOS = ("595", "695")


def normalizar_texto(texto: Any) -> str:
    if not texto:
        return ""
    sin_tildes = "".join(
        c for c in unicodedata.normalize("NFKD", str(texto)) if not unicodedata.combining(c)
    )
    return re.sub(r"[^A-Z0-9]", "", sin_tildes.upper())


def limpiar_valor(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if not s or s in ("#N/A", "#VALUE!", "None", "0", "0.0", "0,0"):
        return None
    return s


def cargar_cod_proyecto_por_estudiante() -> dict[str, str]:
    resultado = {}
    for cod_proyecto in PROYECTOS:
        path = ESTADOS_DIR / f"Listado_de_estudiantes_por_estado_{cod_proyecto}.csv"
        if not path.exists():
            continue
        with path.open(encoding="utf-8-sig") as fh:
            lines = fh.readlines()
        reader = csv.reader(lines[1:])
        next(reader)
        for fila in reader:
            if fila and len(fila) > COL_COD_ESTUDIANTE:
                cod = fila[COL_COD_ESTUDIANTE].strip()
                if cod:
                    resultado[cod] = cod_proyecto
    return resultado


def cargar_directorio_docentes_grupos() -> dict[str, str]:
    grupos_path = SILVER_DIR / "grupos_investigacion.json"
    docente_a_grupo = {}
    if not grupos_path.exists():
        return docente_a_grupo
    try:
        data = json.loads(grupos_path.read_text(encoding="utf-8"))
        for g in data.get("grupos", []):
            sigla = g.get("sigla")
            if not sigla:
                continue
            if g.get("lider"):
                docente_a_grupo[normalizar_texto(g["lider"])] = sigla
            for inte in g.get("integrantes", []):
                docente_a_grupo[normalizar_texto(inte.get("nombre"))] = sigla
    except Exception:
        pass
    return docente_a_grupo


def mapear_grupo_por_director(director: str | None, docente_a_grupo: dict[str, str]) -> str | None:
    if not director:
        return None
    dn = normalizar_texto(director)
    if not dn:
        return None
    if dn in docente_a_grupo:
        return docente_a_grupo[dn]
    for k, sigla in docente_a_grupo.items():
        if dn in k or k in dn:
            return sigla
    return None


def clasificar_etapa(
    has_sust: bool,
    has_jur: bool,
    has_ante: bool,
    caracter: str | None,
    concepto_jur1: str | None,
    concepto_jur2: str | None,
    nota_sust: Any,
) -> str:
    if has_sust:
        car = caracter
        if not car:
            c1 = (concepto_jur1 or "").upper()
            c2 = (concepto_jur2 or "").upper()
            if "MERITORIO" in c1 or "MERITORIO" in c2:
                car = "Meritorio"
            elif "LAUREADO" in c1 or "LAUREADO" in c2:
                car = "Laureado"
            elif nota_sust is not None:
                try:
                    val_nota = float(str(nota_sust).replace(",", "."))
                    if val_nota >= 4.5 or val_nota >= 45:
                        car = "Meritorio"
                    else:
                        car = "Aprobado"
                except Exception:
                    car = "Aprobado"
            else:
                car = "Aprobado"
        return f"Sustentado ({car})"
    if has_jur:
        return "Jurados solicitados, pendiente sustentación"
    if has_ante:
        return "Anteproyecto radicado, pendiente asignación de jurados"
    return "Sin anteproyecto radicado"


def procesar_hoja(
    ws,
    source_label: str,
    sheet_name: str,
    default_mod: str,
    codigos_oficiales: dict[str, str],
    docente_a_grupo: dict[str, str],
) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(c).strip().upper() if c is not None else "" for c in rows[0]]
    h_idx = {h: i for i, h in enumerate(headers) if h}

    def get_f(r: tuple, keys: list[str]) -> Any:
        for k in keys:
            if k in h_idx and h_idx[k] < len(r):
                val = r[h_idx[k]]
                if val is not None and str(val).strip() not in ("", "#N/A", "#VALUE!"):
                    return val
        return None

    registros = []

    for r_idx, r in enumerate(rows[1:], start=2):
        if not any(c is not None and str(c).strip() != "" for c in r):
            continue

        raw_codigo = get_f(r, ["CODIGO", "COD. ESTUDIANTE"])
        codigo = str(raw_codigo).strip() if raw_codigo else None

        titulo = limpiar_valor(get_f(r, ["TITULO DE PROYECTO", "NOMBRE DEL PROYECTO"]))
        director = limpiar_valor(get_f(r, ["DIRECTOR", "DOCENTE DIRECTOR", "DIRECTOR2", "DIRECTOR3"]))
        codirector = limpiar_valor(get_f(r, ["CO-DIRECTOR", "CODIRECTOR"]))
        grupo = limpiar_valor(get_f(r, ["GRUPO DE INVESTIGACION", "GRUPO INVESTIGACION"]))

        # Modalidad / cod_proyecto
        mod_str = str(get_f(r, ["MODALIDAD", "PROYECTO"]) or "").upper()
        if "INVESTIG" in mod_str:
            cod_proyecto = "595"
        elif "PROFUNDIZ" in mod_str or "PASANT" in mod_str:
            cod_proyecto = "695"
        elif default_mod != "ambos":
            cod_proyecto = default_mod
        else:
            enf = str(get_f(r, ["ENFASIS"]) or "")
            if enf in ("595", "695"):
                cod_proyecto = enf
            elif codigo and codigo in codigos_oficiales:
                cod_proyecto = codigos_oficiales[codigo]
            else:
                cod_proyecto = "595"

        # Indicadores de sustentación
        fecha_sust = get_f(r, ["FECHA SUSTENTACION", "FECHA ACTA SUSTENTACION"])
        nota_sust = get_f(r, ["NOTA SUSTENTACION", "NOTA 2 SUSTENTACIÓN", "NOTA 2 SUSTENTACIÓN2", "NOTA"])
        caracter = limpiar_valor(get_f(r, ["CARACTER TESIS"]))
        acta_sust = limpiar_valor(get_f(r, ["ACTA DE SUSTENTACION", "ACTA"]))
        estado_pas = str(get_f(r, ["ESTADO2"]) or "").upper()
        concepto_jur1 = limpiar_valor(get_f(r, ["CONCEPTO JURADO 1"]))
        concepto_jur2 = limpiar_valor(get_f(r, ["CONCEPTO JURADO 2"]))

        has_sust = False
        if fecha_sust or acta_sust or caracter:
            has_sust = True
        elif nota_sust and str(nota_sust).strip() not in ("0", "0.0", "0,0"):
            has_sust = True
        elif "FINALIZADA" in estado_pas:
            has_sust = True
        elif concepto_jur1 and concepto_jur1.upper() in ("APROBADO", "MERITORIO", "LAUREADO"):
            if concepto_jur2 and concepto_jur2.upper() in ("APROBADO", "MERITORIO", "LAUREADO"):
                has_sust = True

        # Indicadores de jurado
        acta_jur = get_f(r, ["ACTA ASIGNACIÓN JURADOS"])
        fecha_jur = get_f(r, ["FECHA ASIGNACION JURADOS", "FECHA RECIBIDO PROY.  JURADO 1"])
        jurado1 = limpiar_valor(get_f(r, ["JURADO 1", "JURADO 1  ACTUAL"]))
        jurado2 = limpiar_valor(get_f(r, ["JURADO 2", "JURADO 2  ACTUAL"]))
        radic_final = get_f(r, ["FECHA DE RADICACION PROYECO FINAL", "FECHA RADICACION PROYECTO FINAL", "FECHA DE RADICACION"])

        has_jur = False
        if not has_sust:
            if acta_jur or fecha_jur or jurado1 or jurado2 or radic_final:
                has_jur = True
            elif (concepto_jur1 and "LISTO" in concepto_jur1.upper()) or (concepto_jur2 and "LISTO" in concepto_jur2.upper()):
                has_jur = True

        # Indicadores de anteproyecto
        fecha_antep = get_f(r, ["FECHA DE RADICACION ANTEPROYECTO", "FECHA RAD. ANTEPROY."])
        acta_antep = get_f(r, ["ACTA ACEPTACION PROPUESTA", "ACTA ACTUAL", "ACTA ACEP. PROPUESTA"])
        fecha_acept_grp = get_f(r, ["FECHA DE ACEPTACION DEL GRUP. INV.", "FECHA ACEPTACION DEL GRUP. INV."])
        revisores = get_f(r, ["REVISOR 1 ANTEP. ACTUAL", "REVISOR 2 ANTEP. ACTUAL"])

        has_ante = False
        if not has_sust and not has_jur:
            if fecha_antep or acta_antep or fecha_acept_grp or revisores:
                has_ante = True
            elif titulo:
                has_ante = True

        # Solo retenemos estudiantes que han iniciado un proceso de trabajo de grado
        if not (has_sust or has_jur or has_ante):
            continue

        etapa = clasificar_etapa(
            has_sust=has_sust,
            has_jur=has_jur,
            has_ante=has_ante,
            caracter=caracter,
            concepto_jur1=concepto_jur1,
            concepto_jur2=concepto_jur2,
            nota_sust=nota_sust,
        )

        # Si el grupo no está explícito (p.ej. Acuerdo Antiguo), asociar por director
        if not grupo and director:
            grupo = mapear_grupo_por_director(director, docente_a_grupo)

        jurados_list = [j for j in [jurado1, jurado2] if j]

        registros.append({
            "_codigo": codigo,
            "_source_priority": 2 if ("INVESTIGACION" in source_label or "Profundizacion" in source_label) else 1,
            "cod_proyecto": cod_proyecto,
            "etapa_actual": etapa,
            "titulo_proyecto": titulo,
            "director": director,
            "codirector": codirector,
            "grupo_investigacion": grupo,
            "numero_acta_sustentacion": str(acta_sust) if acta_sust else None,
            "nota_sustentacion": str(nota_sust) if nota_sust else None,
            "caracter_sustentacion": caracter if has_sust else None,
            "jurados_sustentacion": jurados_list if jurados_list else None,
            "fuente_categorias": [source_label, sheet_name],
        })

    return registros


def extraer_desde_excel(
    path: Path,
    source_label: str,
    hojas_config: list[tuple[str, str]],
    codigos_oficiales: dict[str, str],
    docente_a_grupo: dict[str, str],
) -> list[dict]:
    if not path.exists():
        print(f"  [AVISO] No se encontró {path.name}")
        return []

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    registros = []
    for sheet_name, default_mod in hojas_config:
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            hoja_regs = procesar_hoja(
                ws,
                source_label=source_label,
                sheet_name=sheet_name,
                default_mod=default_mod,
                codigos_oficiales=codigos_oficiales,
                docente_a_grupo=docente_a_grupo,
            )
            registros.extend(hoja_regs)
    wb.close()
    return registros


def main() -> None:
    codigos_oficiales = cargar_cod_proyecto_por_estudiante()
    docente_a_grupo = cargar_directorio_docentes_grupos()

    todos_los_registros = []

    # 1. Base de datos INVESTIGACION (Cohortes recientes 595)
    todos_los_registros.extend(
        extraer_desde_excel(
            ARCHIVO_INV,
            "MCIC - Base de datos INVESTIGACION.xlsx",
            [("N-A Investigación", "595")],
            codigos_oficiales,
            docente_a_grupo,
        )
    )

    # 2. Base de datos PROFUNDIZACION (Cohortes recientes 695)
    todos_los_registros.extend(
        extraer_desde_excel(
            ARCHIVO_PROF,
            "MCIC - Base de datos Profundizacion.xlsx",
            [("N-A Profundizacion", "695"), ("Pasantías", "695")],
            codigos_oficiales,
            docente_a_grupo,
        )
    )

    # 3. Base de datos V2 (Cohortes históricas e intermedias)
    todos_los_registros.extend(
        extraer_desde_excel(
            ARCHIVO_V2,
            "MCIC - Base de datos V2.xlsx",
            [
                ("N-A Investigación", "595"),
                ("N-A Profundizacion", "695"),
                ("Pasantías", "695"),
                ("Acuerdo Antiguo", "ambos"),
            ],
            codigos_oficiales,
            docente_a_grupo,
        )
    )

    # Desduplicación por código de estudiante
    # Se prioriza la base de datos más reciente (INVESTIGACION.xlsx y Profundizacion.xlsx) sobre V2.xlsx
    proyectos_por_clave: dict[str, dict] = {}
    sin_codigo_idx = 0

    for reg in todos_los_registros:
        cod = reg.get("_codigo")
        if cod:
            clave = cod
        else:
            sin_codigo_idx += 1
            clave = f"sin_cod_{sin_codigo_idx}"

        if clave in proyectos_por_clave:
            existente = proyectos_por_clave[clave]
            if reg["_source_priority"] > existente["_source_priority"]:
                proyectos_por_clave[clave] = reg
        else:
            proyectos_por_clave[clave] = reg

    # Construir listado público sin PII
    procesos = []
    for p in proyectos_por_clave.values():
        entrada_publica = {
            "cod_proyecto": p["cod_proyecto"],
            "etapa_actual": p["etapa_actual"],
            "titulo_proyecto": p["titulo_proyecto"],
            "director": p["director"],
            "codirector": p["codirector"],
            "grupo_investigacion": p["grupo_investigacion"],
            "numero_acta_sustentacion": p["numero_acta_sustentacion"],
            "nota_sustentacion": p["nota_sustentacion"],
            "caracter_sustentacion": p["caracter_sustentacion"],
            "jurados_sustentacion": p["jurados_sustentacion"],
            "fuente_categorias": p["fuente_categorias"],
        }
        procesos.append(entrada_publica)

    # Resumen por etapa global y por modalidad
    resumen_etapas_global: dict[str, int] = {}
    resumen_etapas_inv: dict[str, int] = {}
    resumen_etapas_prof: dict[str, int] = {}

    for p in procesos:
        et = p["etapa_actual"]
        resumen_etapas_global[et] = resumen_etapas_global.get(et, 0) + 1
        if p["cod_proyecto"] == "595":
            resumen_etapas_inv[et] = resumen_etapas_inv.get(et, 0) + 1
        elif p["cod_proyecto"] == "695":
            resumen_etapas_prof[et] = resumen_etapas_prof.get(et, 0) + 1

    data = {
        "procesos": procesos,
        "resumen": {
            "total_procesos": len(procesos),
            "por_etapa": resumen_etapas_global,
            "por_modalidad": {
                "investigacion": {
                    "total": sum(1 for p in procesos if p["cod_proyecto"] == "595"),
                    "por_etapa": resumen_etapas_inv,
                },
                "profundizacion": {
                    "total": sum(1 for p in procesos if p["cod_proyecto"] == "695"),
                    "por_etapa": resumen_etapas_prof,
                },
            },
            "fuentes_consultadas": [
                "MCIC - Base de datos INVESTIGACION.xlsx",
                "MCIC - Base de datos Profundizacion.xlsx",
                "MCIC - Base de datos V2.xlsx",
            ],
        },
        "metodologia": (
            "Datos consolidados a partir de las tres bases de datos oficiales de seguimiento de trabajo "
            "de grado de la Maestría en Ciencias de la Información y las Comunicaciones (MCIC): "
            "MCIC - Base de datos INVESTIGACION.xlsx, MCIC - Base de datos Profundizacion.xlsx y "
            "MCIC - Base de datos V2.xlsx. Se extrae cada proceso de trabajo de grado de las modalidades de "
            "Investigación (595) y Profundización (695), clasificando la etapa en: Sustentado (Aprobado / "
            "Meritorio / Laureado), Jurados solicitados (en evaluación / pendiente sustentación) y "
            "Anteproyecto radicado (pendiente jurados). No se publica el código ni el nombre del estudiante "
            "para garantizar la estricta confidencialidad de los datos personales."
        ),
    }

    SILVER_DIR.mkdir(parents=True, exist_ok=True)
    out_path = SILVER_DIR / "seguimiento_tesis.json"
    out_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Seguimiento de tesis: {len(procesos)} proyectos extraídos -> {out_path.relative_to(ROOT)}")
    print(f"  - Investigación (595): {data['resumen']['por_modalidad']['investigacion']['total']}")
    print(f"  - Profundización (695): {data['resumen']['por_modalidad']['profundizacion']['total']}")
    print("  - Distribución global por etapa:")
    for etapa, conteo in resumen_etapas_global.items():
        print(f"      * {etapa}: {conteo}")


if __name__ == "__main__":
    main()
