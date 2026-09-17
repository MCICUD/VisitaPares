"""Bronze -> Silver

Lee los "Cuadros Maestros" oficiales del CNA (Consejo Nacional de
Acreditación) que la coordinación ya diligenció para el proceso de
renovación de Acreditación de Alta Calidad, y que hoy el sitio no usa en
absoluto — son evidencia real para los Factores 4 (Egresados), 8 (Aportes
de la investigación) y 9 (Bienestar) del Plan de Mejoramiento.

Hay dos archivos, uno por modalidad:
    Data/Bronze/ACREDITACIÓN DE ALTA CALIDAD/SNIES17528-MCIC-Investigación/
        SNIES17528-Anexos espc investigación/CuadroMaestro_AcreditacionPrograma Investigaicion.xlsx
    Data/Bronze/ACREDITACIÓN DE ALTA CALIDAD/SNIES116070-MCIC-Profundización/
        SNIES116070-Anexos esp Profund/CuadroMaestro_AcreditacionPrograma Profundizacion.xlsx

Las hojas "Graduación", "Estudiantes", "Estadisticas Bienestar" e
"Investigacion - grupos y profe" resultaron ser LA MISMA TABLA en ambos
archivos (verificado celda a celda): reportan cifras del programa MCIC
completo (matrícula, graduación, bienestar, grupos de investigación), no
desagregadas por modalidad Investigación/Profundización. Por eso aquí se
leen una sola vez, del archivo de Investigación, y se cita como fuente
única — leerlas dos veces solo duplicaría el mismo dato con una ruta
distinta, no agregaría información nueva. (La única diferencia real
encontrada entre ambos archivos en la hoja "Estudiantes" es un error de
fórmula "#DIV/0!" en la copia de Profundización donde Investigación tiene
la celda vacía; no es un dato distinto.)
"""
from __future__ import annotations

import json
import sys
import warnings
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(Path(__file__).resolve().parent))

from lib.xlsx_reader import fuente  # noqa: E402

SILVER_DIR = ROOT / "Data/Silver"
ARCHIVO = (
    ROOT
    / "Data/Bronze/ACREDITACIÓN DE ALTA CALIDAD/SNIES17528-MCIC-Investigación"
    / "SNIES17528-Anexos espc investigación/CuadroMaestro_AcreditacionPrograma Investigaicion.xlsx"
)


def rel(path: Path) -> str:
    return str(path.relative_to(ROOT)).replace("\\", "/")


def leer_graduacion(ws) -> dict:
    """Cuadro CNA No. 04: matrícula y graduados por periodo académico."""
    filas = list(ws.iter_rows(values_only=True))
    por_periodo = []
    anio_actual = None
    fila_inicio = None
    fila_fin = None
    for i, fila in enumerate(filas[5:], start=6):  # fila 6 en adelante (1-indexado): primer periodo real
        if fila[0] == "Promedio" or fila[0] == "Total" or fila[0] is None and fila[1] is None:
            break
        if fila[0] is not None:
            anio_actual = fila[0]
        if fila[1] not in ("I", "II"):
            continue
        if fila_inicio is None:
            fila_inicio = i
        fila_fin = i
        por_periodo.append({
            "periodo": f"{anio_actual}-{fila[1]}",
            "matriculados": fila[2],
            "graduados": fila[3],
        })
    return {
        "por_periodo": por_periodo,
        "fuente": fuente(rel(ARCHIVO), "Graduación", f"{fila_inicio}-{fila_fin}"),
    }


def leer_estudiantes_graduados(ws) -> dict:
    """Cuadro CNA No. 02, columna 'Nº total de Graduados' — se usa para
    verificar que coincide con la hoja 'Graduación' (Cuadro 04)."""
    filas = list(ws.iter_rows(values_only=True))
    por_periodo = {}
    anio_actual = None
    fila_inicio = None
    fila_fin = None
    for i, fila in enumerate(filas[7:], start=8):
        if fila[0] == "Promedio":
            break
        if fila[0] is not None:
            anio_actual = fila[0]
        if fila[1] not in ("I", "II"):
            continue
        if fila_inicio is None:
            fila_inicio = i
        fila_fin = i
        por_periodo[f"{anio_actual}-{fila[1]}"] = fila[8]
    return {
        "por_periodo": por_periodo,
        "fuente": fuente(rel(ARCHIVO), "Estudiantes", f"{fila_inicio}-{fila_fin}"),
    }


def leer_bienestar(ws) -> dict:
    """Cuadro CNA No. 10: actividades y estudiantes atendidos por servicio
    de bienestar, por semestre."""
    filas = list(ws.iter_rows(values_only=True))
    header_semestres = filas[8]  # fila 9 (1-indexada): etiqueta de semestre cada 2 columnas, desde la col E (índice 4)
    semestres = []
    for col in range(4, 24, 2):
        semestres.append((header_semestres[col], col))

    servicios = []
    fila_inicio = None
    fila_fin = None
    for i, fila in enumerate(filas[10:], start=11):
        nombre = fila[3]
        if nombre is None or str(nombre).strip().lower() == "totales":
            break
        if fila_inicio is None:
            fila_inicio = i
        fila_fin = i
        por_semestre = {}
        for semestre, col in semestres:
            por_semestre[semestre] = {
                "actividades": fila[col],
                "estudiantes_atendidos": fila[col + 1],
            }
        servicios.append({"servicio": nombre, "por_semestre": por_semestre})

    return {
        "semestres": [s for s, _ in semestres],
        "servicios": servicios,
        "fuente": fuente(rel(ARCHIVO), "Estadisticas Bienestar", f"{fila_inicio}-{fila_fin}"),
    }


def leer_grupos(ws) -> dict:
    """Cuadro CNA No. 08: producción investigativa por grupo."""
    filas = list(ws.iter_rows(values_only=True))
    grupos = []
    fila_inicio = None
    fila_fin = None
    for i, fila in enumerate(filas[5:], start=6):
        nombre = fila[2]
        if nombre is None:
            break
        if fila_inicio is None:
            fila_inicio = i
        fila_fin = i
        grupos.append({
            "nombre_grupo": nombre,
            "lineas_investigacion": fila[1],
            "codigo_minciencias": fila[3],
            "clasificacion_minciencias": fila[4],
            "proyectos_recursos_internos": fila[5],
            "proyectos_recursos_externos": fila[6],
            "articulos_indexados_nacional": fila[7],
            "articulos_indexados_internacional": fila[8],
            "libros": fila[9],
            "patentes": fila[10],
            "productos_creacion": fila[11],
            "productos_totales": fila[12],
        })
    return {
        "grupos": grupos,
        "fuente": fuente(rel(ARCHIVO), "Investigacion - grupos y profe", f"{fila_inicio}-{fila_fin}"),
    }


def main() -> None:
    if not ARCHIVO.exists():
        raise FileNotFoundError(f"No existe {ARCHIVO}")

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = openpyxl.load_workbook(ARCHIVO, data_only=True, read_only=True)

    graduacion = leer_graduacion(wb["Graduación"])
    graduados_cuadro_estudiantes = leer_estudiantes_graduados(wb["Estudiantes"])

    discrepancias_graduacion = []
    por_periodo_04 = {p["periodo"]: p["graduados"] for p in graduacion["por_periodo"]}
    for periodo, graduados_02 in graduados_cuadro_estudiantes["por_periodo"].items():
        graduados_04 = por_periodo_04.get(periodo)
        if graduados_04 != graduados_02:
            discrepancias_graduacion.append({
                "periodo": periodo,
                "graduados_cuadro_04": graduados_04,
                "graduados_cuadro_02": graduados_02,
            })
    graduados_cuadro_estudiantes["discrepancias_vs_cuadro_04"] = discrepancias_graduacion

    bienestar = leer_bienestar(wb["Estadisticas Bienestar"])
    grupos_produccion = leer_grupos(wb["Investigacion - grupos y profe"])

    data = {
        "nota_alcance": (
            "Estas 4 hojas del Cuadro Maestro CNA reportan cifras del programa MCIC completo (ambas "
            "modalidades juntas), no desagregadas por Investigación/Profundización — se verificó que el "
            "archivo de cada modalidad trae la misma tabla; aquí se cita solo la copia de Investigación."
        ),
        "graduacion": graduacion,
        "graduados_cuadro_estudiantes": graduados_cuadro_estudiantes,
        "bienestar": bienestar,
        "grupos_produccion": grupos_produccion,
    }

    SILVER_DIR.mkdir(parents=True, exist_ok=True)
    out_path = SILVER_DIR / "cuadros_maestros.json"
    out_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Cuadros Maestros CNA extraídos -> {out_path.relative_to(ROOT)}")
    print(f"  - Graduación: {len(graduacion['por_periodo'])} periodos (2019-I a 2024-II)")
    print(f"  - Discrepancias Cuadro 04 vs Cuadro 02: {len(discrepancias_graduacion)}")
    print(f"  - Bienestar: {len(bienestar['servicios'])} servicios x {len(bienestar['semestres'])} semestres")
    print(f"  - Grupos de investigación con producción: {len(grupos_produccion['grupos'])}")


if __name__ == "__main__":
    main()
