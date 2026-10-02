"""Bronze -> Silver

Autoevaluación 2025 (anterior) y 2026 (vigente) de cada modalidad, para el
detalle del Factor 11: documentos de cada proceso (carpetas «Autoevaluación
2025 (anterior)» y «Autoevaluación 2026 (vigente)»), cuántas personas
respondieron cada instrumento y el promedio por factor de 2026, leídos del
libro F11_Autoevaluacion_2025_vs_2026_<modalidad>.xlsx que acompaña a esos
documentos.
"""
from __future__ import annotations

import json
import re
import unicodedata
import warnings
from pathlib import Path

import openpyxl

from build_bronze_manifest import tamano_legible

ROOT = Path(__file__).resolve().parent.parent
SILVER_DIR = ROOT / "Data/Silver"
MODALIDADES = {
    "investigacion": ROOT / "Data/Bronze/MCIC.INVESTIGACION/Procesos de Renocavion y acreditación/Plan de Mejoramiento",
    "profundizacion": ROOT / "Data/Bronze/MCIC-PROFUNDIZACION/Procesos de Renocavion y acreditación/Plan de Mejoramiento",
}
NOMBRES_FACTOR = {1: "Proyecto educativo", 2: "Estudiantes", 3: "Profesores", 4: "Egresados", 5: "Aspectos académicos y resultados de aprendizaje",
                  6: "Permanencia y graduación", 7: "Interacción con el entorno", 8: "Aportes de la investigación", 9: "Bienestar",
                  10: "Medios educativos y ambientes de aprendizaje", 11: "Organización, administración y financiación", 12: "Recursos físicos y tecnológicos"}


def rel(p: Path) -> str:
    return str(p.relative_to(ROOT)).replace("\\", "/")


def nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s)


def documentos(factor_dir: Path, anio: str) -> tuple[list[dict], int]:
    """(documentos principales, total de anexos) de las carpetas del año dentro del Factor 11."""
    principales, anexos = [], 0
    for f in sorted(p for p in factor_dir.rglob("*") if p.is_file()):
        partes = [nfc(x) for x in f.relative_to(factor_dir).parts]
        if not any(anio in x and "Autoevaluación" in x for x in partes):
            continue
        if any(x.lower() == "anexos" for x in partes[:-1]):
            anexos += 1
            continue
        principales.append({"nombre": f.name, "archivo": rel(f), "tamano_legible": tamano_legible(f.stat().st_size)})
    return principales, anexos


def leer_libro(path: Path) -> dict:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    resp = {}
    for fila in wb["Resumen"].iter_rows(values_only=True):
        pass
    celdas = [[c for c in r] for r in wb["Resumen"].iter_rows(values_only=True)]
    for i, fila in enumerate(celdas[:-1]):
        for j, c in enumerate(fila):
            m = re.match(r"(Estudiantes|Docentes|Egresados|Directivos) que respondieron (2025|2026-1)", str(c or ""))
            if m and celdas[i + 1][j] not in (None, "—"):
                resp.setdefault("2025" if m.group(2) == "2025" else "2026", {})[m.group(1)] = celdas[i + 1][j]
    comparacion = [{"aspecto": r[0], "anio_2025": r[1], "anio_2026": r[2]} for r in wb["2025 vs 2026"].iter_rows(min_row=4, values_only=True) if r[0]]
    promedios = []
    for r in wb["Resultados 2026-1"].iter_rows(min_row=5, values_only=True):
        if r[0] and r[1] and str(r[1]).startswith("FACTOR"):
            n = int(str(r[1]).split()[1])
            promedios.append({"instrumento": r[0], "factor": n, "factor_nombre": NOMBRES_FACTOR.get(n, ""), "respuestas": r[2], "preguntas": r[3], "promedio": r[4]})
    return {"respondentes": resp, "comparacion": comparacion, "promedios": promedios}


def main() -> None:
    SILVER_DIR.mkdir(parents=True, exist_ok=True)
    out = {}
    for modalidad, base in MODALIDADES.items():
        factor_dir = next(p for p in base.iterdir() if p.is_dir() and re.match(r"FACTOR 11\.", nfc(p.name)))
        libro = next(factor_dir.rglob("F11_Autoevaluacion_2025_vs_2026_*.xlsx"))
        datos = leer_libro(libro)
        docs25, an25 = documentos(factor_dir, "2025")
        docs26, an26 = documentos(factor_dir, "2026")
        out[modalidad] = {
            "libro": {"nombre": libro.name, "archivo": rel(libro), "tamano_legible": tamano_legible(libro.stat().st_size)},
            "anio_2025": {"titulo": "Autoevaluación 2025 (anterior)", "respondentes": datos["respondentes"].get("2025", {}), "documentos": docs25, "anexos_total": an25},
            "anio_2026": {"titulo": "Autoevaluación 2026 (vigente)", "respondentes": datos["respondentes"].get("2026", {}), "documentos": docs26, "anexos_total": an26,
                          "promedios_por_factor": datos["promedios"]},
            "comparacion": datos["comparacion"],
        }
    (SILVER_DIR / "autoevaluaciones.json").write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    for m, d in out.items():
        print(f"Autoevaluaciones {m}: 2025={d['anio_2025']['respondentes']} 2026={d['anio_2026']['respondentes']} -> Data/Silver/autoevaluaciones.json")


if __name__ == "__main__":
    main()
