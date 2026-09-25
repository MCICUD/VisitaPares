"""Genera las versiones con data específica de la Maestría (MCIC)
para los documentos entregados en las Solicitudes de Pares del DIA 1.
"""
from __future__ import annotations

import re
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
BASE_DIR = ROOT / "SolicitudesPares/DIA 1/SOLICITUDES DE PARES"

MCIC_REGEX = re.compile(
    r"CIENCIAS DE LA INFORMACI[OÓ]N|MCIC|SNIES\s*17528|SNIES\s*116070",
    re.IGNORECASE,
)


def delete_rows_in_batches(ws, rows_to_delete: list[int]) -> None:
    """Elimina filas agrupando índices contiguos para alta velocidad."""
    if not rows_to_delete:
        return
    rows = sorted(set(rows_to_delete))
    spans = []
    start = rows[0]
    count = 1
    for r in rows[1:]:
        if r == start + count:
            count += 1
        else:
            spans.append((start, count))
            start = r
            count = 1
    spans.append((start, count))

    for start, count in reversed(spans):
        ws.delete_rows(start, count)


def filter_docentes_planta() -> None:
    src = BASE_DIR / "Docentes/DOCENTES DE PLANTA GENERAL 2026-1.xlsx"
    dst = BASE_DIR / "Docentes/DOCENTES DE PLANTA GENERAL 2026-1 (MCIC Filtrado).xlsx"
    print(f"Filtrando {src.name}...")
    wb = openpyxl.load_workbook(src)
    ws = wb["PLANTA"]

    rows_to_delete = []
    for r in range(9, ws.max_row + 1):
        proy = str(ws.cell(r, 4).value or "")
        if not MCIC_REGEX.search(proy):
            rows_to_delete.append(r)

    delete_rows_in_batches(ws, rows_to_delete)

    for idx, r in enumerate(range(9, ws.max_row + 1), start=1):
        ws.cell(r, 2).value = idx

    wb.save(dst)
    print(f"  -> Guardado {dst.name} con {ws.max_row - 8} docentes MCIC.")


def filter_biblioteca_proyectos() -> None:
    src = BASE_DIR / "Biblioteca/Biblioteca Estadísticas x Proyecto Curricular 2026-1.xlsx"
    dst = BASE_DIR / "Biblioteca/Biblioteca Estadísticas x Proyecto Curricular 2026-1 (MCIC Filtrado).xlsx"
    print(f"Filtrando {src.name}...")
    wb = openpyxl.load_workbook(src)

    for sname in wb.sheetnames:
        if sname == "Contenido":
            continue
        ws = wb[sname]
        header_r = 1
        for r in range(1, min(10, ws.max_row + 1)):
            row_txt = " ".join(str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1)).upper()
            if "PROYECTO" in row_txt or "PROGRAMA" in row_txt or "FACULTAD" in row_txt:
                header_r = r
                break

        rows_to_delete = []
        for r in range(header_r + 1, ws.max_row + 1):
            row_txt = " ".join(str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1))
            if not MCIC_REGEX.search(row_txt):
                rows_to_delete.append(r)

        delete_rows_in_batches(ws, rows_to_delete)

    wb.save(dst)
    print(f"  -> Guardado {dst.name}")


def filter_biblioteca_inversion() -> None:
    src = BASE_DIR / "Biblioteca/Biblioteca Inversión_Adquisición Recusos Bibliograficos 2026-1.xlsx"
    dst = BASE_DIR / "Biblioteca/Biblioteca Inversión_Adquisición Recusos Bibliograficos 2026-1 (MCIC Filtrado).xlsx"
    print(f"Filtrando {src.name}...")
    wb = openpyxl.load_workbook(src)

    for sname in wb.sheetnames:
        ws = wb[sname]
        header_r = 1
        for r in range(1, min(10, ws.max_row + 1)):
            row_txt = " ".join(str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1)).upper()
            if "PROYECTO" in row_txt or "PROGRAMA" in row_txt or "SNIES" in row_txt:
                header_r = r
                break

        rows_to_delete = []
        for r in range(header_r + 1, ws.max_row + 1):
            row_txt = " ".join(str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1))
            if not MCIC_REGEX.search(row_txt):
                rows_to_delete.append(r)

        delete_rows_in_batches(ws, rows_to_delete)

    wb.save(dst)
    print(f"  -> Guardado {dst.name}")


def filter_investigacion_grupos() -> None:
    src = BASE_DIR / "Investigación/Grupos-Investig(957).xlsx"
    dst = BASE_DIR / "Investigación/Grupos-Investig(957) (MCIC Filtrado).xlsx"
    print(f"Filtrando {src.name}...")
    wb = openpyxl.load_workbook(src)

    # Filtrar Docentes 2025-3
    ws_doc = wb["Docentes 2025-3"]
    rows_to_delete_doc = []
    for r in range(2, ws_doc.max_row + 1):
        txt = " ".join(str(ws_doc.cell(r, c).value or "") for c in range(1, ws_doc.max_column + 1))
        if not MCIC_REGEX.search(txt):
            rows_to_delete_doc.append(r)
    delete_rows_in_batches(ws_doc, rows_to_delete_doc)

    # Filtrar GruposUD 2025-3 para grupos vinculados a Ingeniería / MCIC
    ws_grp = wb["GruposUD 2025-3"]
    rows_to_delete_grp = []
    for r in range(2, ws_grp.max_row + 1):
        txt = " ".join(str(ws_grp.cell(r, c).value or "") for c in range(1, ws_grp.max_column + 1)).upper()
        if "FACULTAD DE INGENIERÍA" not in txt and "FACULTAD DE INGENIERIA" not in txt:
            rows_to_delete_grp.append(r)
    delete_rows_in_batches(ws_grp, rows_to_delete_grp)

    wb.save(dst)
    print(f"  -> Guardado {dst.name}")


def filter_oferta_academica() -> None:
    src = BASE_DIR / "Programas registro calificado y acreditación/OFERTA ACADÉMICA 14-08-2026.xlsx"
    dst = BASE_DIR / "Programas registro calificado y acreditación/OFERTA ACADÉMICA 14-08-2026 (MCIC Filtrado).xlsx"
    print(f"Filtrando {src.name}...")
    wb = openpyxl.load_workbook(src)
    ws = wb["Programas en oferta"]

    # Fila de encabezado es la 6; datos de programas van del 7 al 106.
    header_r = 6
    rows_to_delete = []
    for r in range(header_r + 1, 107):
        txt = " ".join(str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1))
        if not MCIC_REGEX.search(txt):
            rows_to_delete.append(r)

    # Descombinar celdas en filas a eliminar
    to_delete_set = set(rows_to_delete)
    merged_to_remove = [
        m for m in ws.merged_cells.ranges
        if any(r in to_delete_set for r in range(m.min_row, m.max_row + 1))
    ]
    for m in merged_to_remove:
        ws.merged_cells.remove(m)

    delete_rows_in_batches(ws, rows_to_delete)

    # Renumerar los programas MCIC
    for idx, r in enumerate(range(header_r + 1, header_r + 3), start=1):
        ws.cell(r, 1).value = idx

    wb.save(dst)
    print(f"  -> Guardado {dst.name}")


def filter_seguimiento_aacpc() -> None:
    src = BASE_DIR / "Programas registro calificado y acreditación/Seguimiento AACPC_15_09_26.xlsx"
    dst = BASE_DIR / "Programas registro calificado y acreditación/Seguimiento AACPC_15_09_26 (MCIC Filtrado).xlsx"
    print(f"Filtrando {src.name}...")
    wb = openpyxl.load_workbook(src)

    # Seguimiento Gral: encabezado en fila 1, fila 2 vacía, datos en fila 3..max_row
    ws = wb["Seguimiento Gral"]
    header_r = 2
    rows_to_delete = []
    for r in range(header_r + 1, ws.max_row + 1):
        txt = " ".join(str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1))
        if not MCIC_REGEX.search(txt):
            rows_to_delete.append(r)

    delete_rows_in_batches(ws, rows_to_delete)

    # En las otras hojas (Acreditables Ac 02 2020 y Acreditables Ac 01 2025),
    # filtrar también para dejar solo filas MCIC si las hubiera, o dejarlas limpias
    for sname in ["Acreditables Ac 02 2020", "Acreditables Ac 01 2025"]:
        if sname in wb.sheetnames:
            ws_other = wb[sname]
            other_del = []
            for r in range(2, ws_other.max_row + 1):
                txt = " ".join(str(ws_other.cell(r, c).value or "") for c in range(1, ws_other.max_column + 1))
                if not MCIC_REGEX.search(txt):
                    other_del.append(r)
            delete_rows_in_batches(ws_other, other_del)

    wb.save(dst)
    print(f"  -> Guardado {dst.name}")


def filter_movilidad_estudiantes() -> None:
    src = BASE_DIR / "Unidad de relaciones internacionales/MOVILIDAD ESTUDIANTES ENTRANTE Y SALIENTE 2021-1 2026-1.xlsx"
    dst = BASE_DIR / "Unidad de relaciones internacionales/MOVILIDAD ESTUDIANTES ENTRANTE Y SALIENTE 2021-1 2026-1 (MCIC Filtrado).xlsx"
    print(f"Filtrando {src.name}...")
    wb = openpyxl.load_workbook(src)

    for sname in wb.sheetnames:
        ws = wb[sname]
        header_r = 1
        rows_to_delete = []
        for r in range(header_r + 1, ws.max_row + 1):
            txt = " ".join(str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1))
            if not MCIC_REGEX.search(txt):
                rows_to_delete.append(r)

        delete_rows_in_batches(ws, rows_to_delete)

    wb.save(dst)
    print(f"  -> Guardado {dst.name}")


def main() -> None:
    print("=== Generando versiones con data específica de MCIC ===")
    filter_docentes_planta()
    filter_biblioteca_proyectos()
    filter_biblioteca_inversion()
    filter_investigacion_grupos()
    filter_oferta_academica()
    filter_seguimiento_aacpc()
    filter_movilidad_estudiantes()
    print("=== Filtrado completado con éxito. ===")


if __name__ == "__main__":
    main()
