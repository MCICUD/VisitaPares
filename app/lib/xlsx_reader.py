"""Utilidades mínimas para leer hojas de Excel con openpyxl y conservar,
para cada valor extraído, la referencia exacta (archivo, hoja, fila) que
permite verificarlo contra el documento fuente.
"""
from __future__ import annotations

import datetime
import warnings
from pathlib import Path
from typing import Any

import openpyxl


def load_sheet(path: str | Path, sheet_name: str):
    """Abre un .xlsx y devuelve la hoja pedida (valores calculados, no fórmulas)."""
    with warnings.catch_warnings():
        # openpyxl advierte sobre extensiones de formato condicional / validación
        # de datos que no soporta; no afectan la lectura de valores.
        warnings.simplefilter("ignore", UserWarning)
        wb = openpyxl.load_workbook(path, data_only=True)
    return wb[sheet_name]


def read_table(ws, start_row: int, end_row: int, columns: dict[str, str]) -> list[dict[str, Any]]:
    """Lee filas [start_row, end_row] (inclusive, 1-indexadas) de `ws`.

    `columns` mapea nombre_de_campo -> letra_de_columna (p.ej. {"factor": "B"}).
    Cada fila devuelta incluye `_fila` con el número de fila real en la hoja,
    para poder construir el bloque `fuente` de trazabilidad.
    """
    rows = []
    for row in range(start_row, end_row + 1):
        record: dict[str, Any] = {"_fila": row}
        for field, col_letter in columns.items():
            record[field] = ws[f"{col_letter}{row}"].value
        rows.append(record)
    return rows


def read_cell(ws, coord: str) -> Any:
    return ws[coord].value


def as_flag(value: Any) -> bool:
    """Convierte celdas de marcación ('x'/'X'/None) en booleano."""
    return isinstance(value, str) and value.strip().lower() == "x"


def as_iso_date(value: Any) -> str | None:
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.date().isoformat() if isinstance(value, datetime.datetime) else value.isoformat()
    return value


def fuente(archivo: str, hoja: str, fila: int) -> dict[str, Any]:
    return {"archivo": archivo, "hoja": hoja, "fila": fila}
