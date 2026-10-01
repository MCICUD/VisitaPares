"""Estilo común para todos los .xlsx de NuevaData/Presentacion.

Paleta institucional sobria (azul UD + gris), encabezados fijos, filtros,
bandas alternas y anchos de columna razonables para que todos los libros
se vean iguales.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = "0B3B60"
AZUL_CLARO = "DCE6F1"
GRIS = "F3F4F6"
GRIS_TEXTO = "4B5563"
VERDE = "E8F5E9"
AMBAR = "FFF4E5"

_thin = Side(style="thin", color="C9D3DD")
BORDE = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

FONT_TITULO = Font(name="Calibri", size=15, bold=True, color=AZUL)
FONT_SUB = Font(name="Calibri", size=10.5, italic=True, color=GRIS_TEXTO)
FONT_HEADER = Font(name="Calibri", size=10.5, bold=True, color="FFFFFF")
FONT_CELDA = Font(name="Calibri", size=10)
FONT_LINK = Font(name="Calibri", size=10, color="1F5FA8", underline="single")
FONT_SECCION = Font(name="Calibri", size=11.5, bold=True, color=AZUL)
FONT_KPI = Font(name="Calibri", size=18, bold=True, color=AZUL)

FILL_HEADER = PatternFill("solid", fgColor=AZUL)
FILL_BANDA = PatternFill("solid", fgColor=GRIS)
FILL_KPI = PatternFill("solid", fgColor=AZUL_CLARO)
FILL_NOTA = PatternFill("solid", fgColor=AMBAR)
FILL_OK = PatternFill("solid", fgColor=VERDE)

WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)


def nuevo_libro() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def hoja(wb: Workbook, nombre: str, titulo: str, subtitulo: str | None = None, ancho_titulo: int = 8):
    ws = wb.create_sheet(nombre[:31])
    ws.sheet_view.showGridLines = False
    ws["A1"] = titulo
    ws["A1"].font = FONT_TITULO
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ancho_titulo)
    ws.row_dimensions[1].height = 24
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = FONT_SUB
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ancho_titulo)
        ws.row_dimensions[2].height = max(16, 15 * (1 + len(subtitulo) // 140))
    return ws


def seccion(ws, fila: int, texto: str, ancho: int = 8) -> int:
    ws.cell(row=fila, column=1, value=texto).font = FONT_SECCION
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ancho)
    return fila + 1


def nota(ws, fila: int, texto: str, ancho: int = 8) -> int:
    c = ws.cell(row=fila, column=1, value=texto)
    c.font = Font(name="Calibri", size=9.5, italic=True, color=GRIS_TEXTO)
    c.fill = FILL_NOTA
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ancho)
    ws.row_dimensions[fila].height = max(15, 13 * (1 + len(texto) // 120))
    return fila + 1


def tabla(ws, fila: int, encabezados: list[str], filas: list[list], anchos: list[int] | None = None,
          filtro: bool = True, congelar: bool = True, links: dict[int, int] | None = None) -> int:
    """Escribe una tabla con estilo. `links` = {indice_columna_valor: indice_columna_url}:
    la celda de la columna valor se convierte en hipervínculo a la URL (la columna URL no se escribe)."""
    links = links or {}
    ocultas = set(links.values())
    cols = [i for i in range(len(encabezados)) if i not in ocultas]
    for j, i in enumerate(cols, start=1):
        c = ws.cell(row=fila, column=j, value=encabezados[i])
        c.font, c.fill, c.alignment, c.border = FONT_HEADER, FILL_HEADER, CENTRO, BORDE
    ws.row_dimensions[fila].height = 30
    inicio = fila
    for k, valores in enumerate(filas):
        r = fila + 1 + k
        for j, i in enumerate(cols, start=1):
            v = valores[i] if i < len(valores) else None
            c = ws.cell(row=r, column=j, value=v)
            c.font, c.alignment, c.border = FONT_CELDA, WRAP_TOP, BORDE
            if k % 2 == 1:
                c.fill = FILL_BANDA
            if i in links and valores[links[i]]:
                c.hyperlink = valores[links[i]]
                c.font = FONT_LINK
    fin = fila + len(filas)
    if anchos:
        for j, w in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    if filtro and filas:
        ws.auto_filter.ref = f"A{inicio}:{get_column_letter(len(cols))}{fin}"
    if congelar:
        ws.freeze_panes = ws.cell(row=inicio + 1, column=1)
    return fin + 2


def kpis(ws, fila: int, items: list[tuple[str, object]], por_fila: int = 4) -> int:
    """Tarjetas de indicadores: etiqueta arriba, valor grande abajo (2 columnas por tarjeta)."""
    for n, (etiqueta, valor) in enumerate(items):
        r = fila + (n // por_fila) * 3
        col = 1 + (n % por_fila) * 2
        a = ws.cell(row=r, column=col, value=etiqueta)
        a.font = Font(name="Calibri", size=9.5, bold=True, color=GRIS_TEXTO)
        a.alignment = CENTRO
        a.fill = FILL_KPI
        ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
        b = ws.cell(row=r + 1, column=col, value=valor)
        b.font = FONT_KPI
        b.alignment = CENTRO
        b.fill = FILL_KPI
        ws.merge_cells(start_row=r + 1, start_column=col, end_row=r + 1, end_column=col + 1)
        ws.row_dimensions[r].height = 28
        ws.row_dimensions[r + 1].height = 30
    filas_usadas = ((len(items) - 1) // por_fila + 1) * 3
    return fila + filas_usadas


def miniatura(ws, celda: str, ruta_imagen: Path, ancho_px: int = 260):
    from PIL import Image as PILImage
    import tempfile

    with PILImage.open(ruta_imagen) as im:
        im = im.convert("RGB")
        ratio = ancho_px / im.width
        im = im.resize((ancho_px, max(1, int(im.height * ratio))))
        tmp = Path(tempfile.mkdtemp()) / (ruta_imagen.stem + ".png")
        im.save(tmp)
    img = XLImage(str(tmp))
    ws.add_image(img, celda)
    return img.height
